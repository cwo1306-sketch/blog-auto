"""검토 대기 중인 고유 기사를 엑셀 검토 큐로 내보낸다."""

from __future__ import annotations

from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .db import Database

KST = timezone(timedelta(hours=9))

COLUMNS = [
    ("ID", 8),
    ("선택", 8),            # 사람이 O/X/보류 입력
    ("키워드 그룹", 14),
    ("키워드", 16),
    ("발행 시각(KST)", 18),
    ("매체", 20),
    ("제목", 60),
    ("요약", 80),
    ("중복 수", 8),          # 이 기사로 묶인 중복 기사 수 = 화제성 신호
    ("원문 링크", 50),
    ("네이버 링크", 50),
    ("메모", 30),
]


def export_review_queue(db: Database, output: Path, mark_exported: bool = False) -> int:
    rows = db.pending_unique()
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "검토 큐"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "C2"

    for r, row in enumerate(rows, start=2):
        pub = datetime.fromisoformat(row["pub_date"]).astimezone(KST)
        values = [
            row["id"], "", row["keyword_group"], row["keyword"],
            pub.strftime("%Y-%m-%d %H:%M"), row["source"], row["title"], row["description"],
            db.duplicate_count(row["id"]), row["original_link"], row["link"], "",
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in (7, 8):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c in (10, 11):
            cell = ws.cell(row=r, column=c)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")

    if rows:
        dv = DataValidation(type="list", formula1='"O,X,보류"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"B2:B{len(rows) + 1}")
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    meta = wb.create_sheet("요약")
    stats = db.stats()
    meta.append(["생성 시각(KST)", datetime.now(KST).strftime("%Y-%m-%d %H:%M")])
    meta.append(["검토 대기", stats.get("pending", 0)])
    meta.append(["전체 저장", stats.get("total", 0)])
    for stage, label in ((1, "URL 일치"), (2, "제목 일치"), (3, "제목 유사"), (4, "본문 지문")):
        meta.append([f"{stage}단계 중복 ({label})", stats.get(f"stage_{stage}", 0)])
    meta.column_dimensions["A"].width = 28

    wb.save(output)
    if mark_exported and rows:
        db.mark_exported([row["id"] for row in rows])
    return len(rows)
