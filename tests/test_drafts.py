import copy
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from naver_cardnews import sheet
from naver_cardnews.db import Article, Database
from naver_cardnews.draft import draft_top
from naver_cardnews.normalize import normalize_title
from naver_cardnews.queue import export_queue, import_drafts
from naver_cardnews.section import parse_relative_time, parse_section

FIX = Path(__file__).parent / "fixtures" / "section_105.html"


def test_parse_section_fixture():
    items = parse_section(FIX.read_text(encoding="utf-8"), "105", now=datetime(2026, 10, 21, 3, 0, tzinfo=timezone.utc))
    assert len(items) == 4
    first = items[1]
    assert first["link"].startswith("https://n.news.naver.com/mnews/article/421/")
    assert "SK쉴더스" in first["title"] and first["press"] == "뉴스1" and first["description"]
    assert first["originallink"] == "" and "2026" in first["pubDate"]


def test_parse_relative_time():
    now = datetime(2026, 10, 21, 3, 0, tzinfo=timezone.utc)
    assert parse_relative_time("2시간전", now).hour == 1
    assert parse_relative_time("30분전", now).minute == 30
    assert parse_relative_time("2026.10.19.", now).day == 18  # KST 0시 → UTC 전날 15시


from tests.test_cardnews import FULL, ARTICLE_BODY  # noqa: E402


def test_flatten_unflatten_roundtrip():
    flat = sheet.flatten(FULL)
    assert flat["표지2"] == "*AI의 눈" and flat["3장 유형"] == "stat" and flat["3장 수치"] == "12.5"
    back = sheet.unflatten(flat)
    assert back["cover"]["lines"][1] == {"text": "AI의 눈", "highlight": True}
    assert back["slides"][1]["stat"]["unit"] == "%" and back["slides"][0]["stat"] is None
    assert back["hashtags"] == FULL["hashtags"] and back["caption"] == FULL["caption"]
    assert not sheet.changed(FULL, back)
    edited = dict(flat, **{"2장 제목": "바뀐 제목"})
    assert sheet.changed(FULL, sheet.unflatten(edited))


def _seed(db):
    for i, aid in enumerate(["001_1", "002_2", "003_3"]):
        db.insert(Article(aid, "AI", f"기사 {i}", "a.com", "2026-10-21T00:00:00+00:00",
                          f"https://n.news.naver.com/article/{aid.replace('_', '/')}", "", "요약", normalize_title(f"기사 {i}")),
                  i + 1, True, "대기")
    db.conn.execute("UPDATE articles SET score=3 WHERE id='002_2'")
    db.conn.commit()


def test_draft_top_uses_score_order_and_saves(tmp_path):
    db = Database(tmp_path / "t.db")
    _seed(db)
    calls = []

    def fetcher(url):
        calls.append(url)
        return {"url": url, "id": "x", "title": "t", "press": "p", "date": "2026-10-21", "body": ARTICLE_BODY}

    def write(article, previous=None, errors=None):
        s = copy.deepcopy(FULL); s["_usage"] = {"input_tokens": 100, "output_tokens": 50}
        return s

    rep = draft_top(db, 2, fetcher=fetcher, write=write)
    assert rep.drafted == ["002_2", "001_1"] and rep.input_tokens == 200   # 점수 3 → 먼저
    assert db.get("002_2")["body"] == ARTICLE_BODY
    d = db.latest_draft("002_2")
    assert d["version"] == 1 and d["source"] == "ai" and json.loads(d["errors_json"]) == []
    # 두 번째 실행은 원고 없는 003_3 만
    rep2 = draft_top(db, 5, fetcher=fetcher, write=write)
    assert rep2.drafted == ["003_3"]


def test_export_and_import_edited_draft(tmp_path):
    db = Database(tmp_path / "t.db")
    _seed(db)
    db.set_body("002_2", ARTICLE_BODY)
    db.save_draft("002_2", FULL, [], "ai", {})
    out = tmp_path / "q.xlsx"
    export_queue(db, out)
    wb = load_workbook(out)
    assert {"뉴스큐", "원고", "본문", "요약"} <= set(wb.sheetnames)
    ws = wb["원고"]
    header = [c.value for c in ws[1]]
    assert ws.cell(2, header.index("id") + 1).value == "002_2"
    assert ws.cell(2, header.index("3장 수치") + 1).value == "12.5"
    assert wb["본문"].cell(2, 6).value == ARTICLE_BODY
    # 사람이 2장 본문을 원문에 없는 수치로 고침 → 새 버전 + 검수 오류
    ws.cell(2, header.index("2장 본문") + 1).value = "카카오가 비전 인코더 5종을 공개했습니다. " + "설명을 덧붙입니다. " * 4
    ws.cell(2, header.index("메모") + 1).value = "표지는 숫자 강조"
    wb.save(out)
    assert import_drafts(db, out) == 1
    d = db.latest_draft("002_2")
    assert d["version"] == 2 and d["source"] == "excel"
    assert any("5종" in e for e in json.loads(d["errors_json"]))
    assert db.get("002_2")["memo"] == "표지는 숫자 강조"
    # 변경 없이 다시 반입하면 버전 그대로
    assert import_drafts(db, out) == 0
    export_queue(db, out)   # 오류가 시트에 표시되는지
    ws = load_workbook(out)["원고"]
    header = [c.value for c in ws[1]]
    assert "5종" in ws.cell(2, header.index("검수오류") + 1).value and ws.cell(2, header.index("원고출처") + 1).value == "excel"
