from datetime import datetime, timedelta, timezone
from email.utils import format_datetime
from pathlib import Path

from openpyxl import load_workbook

from naver_cardnews.collect import collect
from naver_cardnews.config import Settings
from naver_cardnews.db import Database
from naver_cardnews.queue import ExcelLocked, export_queue, import_queue


class FakeClient:
    def __init__(self, data):
        self.data = data

    def iter_items(self, query, limit):
        return iter(self.data.get(query, [])[:limit])


def item(oid, aid, title, desc, pub, original="https://press.com/x", naver=True):
    link = f"https://n.news.naver.com/mnews/article/{oid}/{aid}" if naver else original
    return {"title": title, "description": desc, "link": link, "originallink": original,
            "pubDate": format_datetime(pub)}


def test_collect_import_export(tmp_path: Path):
    now = datetime.now(timezone.utc)
    data = {
        "엔비디아": [
            item("015", "1", "<b>엔비디아</b>, 신형 AI 칩 공개", "블랙웰 후속작 루빈을 공개했다", now, "https://hankyung.com/1"),
            item("001", "2", "[속보] 엔비디아 신형 AI 칩 공개", "블랙웰 후속작 루빈을 공개했다", now, "https://yna.co.kr/2"),
            item("003", "9", "오래된 기사", "", now - timedelta(days=10)),
            item("000", "0", "네이버 아님", "", now, "https://other.com/0", naver=False),
        ],
        "AI 반도체": [
            item("015", "1", "엔비디아, 신형 AI 칩 공개", "블랙웰 후속작 루빈을 공개했다", now, "https://hankyung.com/1"),
            item("009", "3", "SK하이닉스 HBM4 공급 계약", "내년 상반기 양산에 들어간다", now, "https://mk.co.kr/3"),
        ],
    }
    s = Settings(client_id="x", client_secret="y", db_path=tmp_path / "t.db",
                 output_path=tmp_path / "q.xlsx", keywords=["엔비디아", "AI 반도체"])
    db = Database(s.db_path)
    rep = collect(s, db, client=FakeClient(data))
    assert (rep.fetched, rep.non_naver, rep.too_old, rep.new_clusters) == (6, 1, 1, 2)
    assert rep.by_stage == {1: 1, 2: 1}
    assert db.get("015_1")["keywords_hit"] == "엔비디아,AI 반도체"
    assert db.get("015_1")["score"] == 2

    # 내보내기
    assert export_queue(db, s.output_path) == 3
    ws = load_workbook(s.output_path)["뉴스큐"]
    header = [c.value for c in ws[1]]
    assert header[:4] == ["id", "수집일시", "키워드", "제목"] and "상태" in header
    rows = {r[0]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["015_1"]["상태"] == "대기" and rows["015_1"]["대표여부"] == "대표"
    assert rows["001_2"]["상태"] == "제외" and rows["001_2"]["중복사유"].startswith("중복(제목)")

    # 사람이 승인 → DB 반영
    wb = load_workbook(s.output_path)
    ws = wb["뉴스큐"]
    sc, mc = header.index("상태") + 1, header.index("메모") + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == "015_1":
            ws.cell(r, sc).value = "승인"
            ws.cell(r, mc).value = "표지는 숫자 카드"
    wb.save(s.output_path)
    assert import_queue(db, s.output_path) == 1
    assert db.get("015_1")["status"] == "승인"
    assert db.get("015_1")["memo"] == "표지는 숫자 카드"
    assert import_queue(db, s.output_path) == 0  # 두 번째는 변화 없음

    # 재내보내기 시 백업 생성, 승인 상태 유지
    export_queue(db, s.output_path)
    assert list((tmp_path / "backups").glob("q_*.xlsx"))
    ws = load_workbook(s.output_path)["뉴스큐"]
    rows = {r[0]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["015_1"]["상태"] == "승인"


def test_excel_lock_detected(tmp_path: Path):
    db = Database(":memory:")
    out = tmp_path / "q.xlsx"
    (tmp_path / "~$q.xlsx").write_text("")
    try:
        export_queue(db, out)
        assert False, "lock not detected"
    except ExcelLocked:
        pass
