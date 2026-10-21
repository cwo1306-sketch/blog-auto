import copy
import json
from pathlib import Path

from openpyxl import load_workbook

from naver_cardnews import feedback as fb
from naver_cardnews import rubric, writer
from naver_cardnews.db import Article, Database
from naver_cardnews.normalize import normalize_title
from naver_cardnews.queue import export_queue, import_scores
from tests.test_cardnews import ARTICLE_BODY, FULL


def test_parse_evaluation_and_total():
    ev = rubric.parse_evaluation('설명 {"scores":{"fact":5,"attribution":3,"plain":4,"structure":5,"hook":2},"issues":[],"unsupported_claims":["x"],"summary":"ok"}')
    assert ev.scores["hook"] == 2 and ev.unsupported_claims == ["x"]
    assert ev.total == round(5*.35 + 3*.15 + 4*.25 + 5*.1 + 2*.15, 2)
    ev2 = rubric.parse_evaluation('{"scores":{"fact":9,"attribution":0}}')
    assert ev2.scores["fact"] == 5 and ev2.scores["attribution"] == 1


def test_weighted_kappa_and_report():
    assert rubric.weighted_kappa([1, 2, 3, 4, 5], [1, 2, 3, 4, 5]) == 1.0
    k = rubric.weighted_kappa([5, 5, 4, 4, 3, 3], [1, 2, 2, 1, 5, 4])
    assert k is not None and k < 0.2
    assert rubric.weighted_kappa([3], [3]) is None
    pairs = [({a: 4 for a in rubric.AXES}, {a: 3 for a in rubric.AXES})] * 6
    rep = rubric.agreement_report(pairs)
    assert rep["axes"]["fact"]["bias"] == 1.0 and rep["axes"]["fact"]["within1"] == 1.0
    lines = rubric.verdict(rep)
    assert any("후하게" in l for l in lines)


def test_plain_metrics():
    pm = rubric.plain_metrics(FULL)
    assert pm["sentences"] > 5 and pm["avg_words"] > 0 and pm["max_words"] >= pm["avg_words"]


def test_classify_change():
    assert fb.classify_change("2장 본문", "12.5% 높았다", "13% 높았다") == "수치 수정"
    assert fb.classify_change("2장 본문", "a" * 100, "a" * 60) == "축약"
    assert fb.classify_change("2장 본문", "카카오가 공개했습니다.", "카카오가 공개했어요.") == "어투 변경"
    assert fb.classify_change("2장 본문", "카카오가 모델을 공개했습니다.", "카카오가 새 모델을 공개했습니다.") == "표현 다듬기"
    assert fb.classify_change("2장 검색어", "a | b", "c | d") == "사진 지시 수정"
    assert fb.classify_change("2장 본문", "x", "") == "삭제"


def test_diff_aggregate_fewshot_roundtrip(tmp_path):
    db = Database(tmp_path / "t.db")
    db.insert(Article("002_2", "AI", "기사", "a.com", "2026-10-21T00:00:00+00:00", "https://n.news.naver.com/article/002/2", "", "", normalize_title("기사")), 1, True, "대기")
    db.set_body("002_2", ARTICLE_BODY)
    v1 = db.save_draft("002_2", FULL, [], "ai", {})
    edited = copy.deepcopy(FULL)
    edited["slides"][0]["body"] = edited["slides"][0]["body"].replace("공개했습니다", "공개했어요").replace("있습니다", "있어요").replace("합니다", "해요")
    edited["slides"][2]["body"] = edited["slides"][2]["body"][:45] + "."
    v2 = db.save_draft("002_2", edited, [], "excel", {})
    diffs = fb.collect_diffs(db)
    assert len(diffs) == 1 and (diffs[0].from_version, diffs[0].to_version) == (v1, v2)
    kinds = {c["field"]: c["kind"] for c in diffs[0].changes}
    assert kinds["2장 본문"] == "어투 변경" and kinds["4장 본문"] == "축약"
    agg = fb.aggregate(diffs)
    assert agg["articles"] == 1 and agg["cells"] == 2
    assert "수정된 원고 1건" in fb.report_text(agg)
    ex = fb.build_fewshot(diffs)
    assert len(ex) == 2 and all(e["kind"] in ("어투 변경", "축약") for e in ex)
    writer.set_examples(fb.fewshot_prompt(ex))
    assert "편집자가 실제로 고친 예시" in writer.system_prompt() and "공개했어요" in writer.system_prompt()
    writer.set_examples("")


def test_scores_sheet_export_and_human_import(tmp_path):
    db = Database(tmp_path / "t.db")
    db.insert(Article("002_2", "AI", "기사", "a.com", "2026-10-21T00:00:00+00:00", "https://n.news.naver.com/article/002/2", "", "", normalize_title("기사")), 1, True, "대기")
    db.set_body("002_2", ARTICLE_BODY)
    v = db.save_draft("002_2", FULL, [], "ai", {})
    db.save_evaluation("002_2", v, "judge", {a: 4 for a in rubric.AXES}, 4.0, {"unsupported_claims": ["근거 없음 문장"], "summary": "총평"})
    out = tmp_path / "q.xlsx"
    export_queue(db, out)
    wb = load_workbook(out)
    assert {"채점", "피드백"} <= set(wb.sheetnames)
    ws = wb["채점"]
    header = [c.value for c in ws[1]]
    assert ws.cell(2, header.index("판정 사실 일치") + 1).value == 4
    assert ws.cell(2, header.index("근거 없는 주장") + 1).value == "근거 없음 문장"
    # 사람이 점수 입력
    for a, val in zip(rubric.AXES, [3, 4, 2, 5, 3]):
        ws.cell(2, header.index(f"사람 {rubric.AXIS_KO[a]}") + 1).value = val
    ws.cell(2, header.index("사람 메모") + 1).value = "출처 주체 빠짐"
    wb.save(out)
    assert import_scores(db, out) == 1
    h = db.evaluation("002_2", v, "human")
    assert json.loads(h["scores_json"])["plain"] == 2 and json.loads(h["detail_json"])["memo"] == "출처 주체 빠짐"
    assert import_scores(db, out) == 0   # 변화 없음
    pairs = db.judge_human_pairs()
    assert len(pairs) == 1 and pairs[0][0]["fact"] == 4 and pairs[0][1]["fact"] == 3


def test_spearman_and_total_report():
    assert rubric.spearman([1, 2, 3], [1, 2, 3]) == 1.0
    assert rubric.spearman([2.3, 2.5, 2.5], [5, 3, 3]) == -1.0
    assert rubric.spearman([1, 2], [1, 2]) is None
    tr = rubric.total_report([("a", 2.3, 5.0), ("b", 2.5, 3.0), ("c", 2.5, 3.0)])
    assert tr["n"] == 3 and tr["bias"] == -1.23 and tr["largest_gaps"][0]["id"] == "a"


def test_overall_only_human_score_import(tmp_path):
    db = Database(tmp_path / "t.db")
    db.insert(Article("002_2", "AI", "기사", "a.com", "2026-10-21T00:00:00+00:00", "https://n.news.naver.com/article/002/2", "", "", normalize_title("기사")), 1, True, "대기")
    v = db.save_draft("002_2", FULL, [], "ai", {})
    db.save_evaluation("002_2", v, "judge", {a: 2 for a in rubric.AXES}, 2.0, {})
    out = tmp_path / "q.xlsx"
    export_queue(db, out)
    wb = load_workbook(out); ws = wb["채점"]
    header = [c.value for c in ws[1]]
    ws.cell(2, header.index("사람 총점") + 1).value = 5
    wb.save(out)
    assert import_scores(db, out) == 1
    h = db.evaluation("002_2", v, "human")
    assert h["total"] == 5.0 and json.loads(h["scores_json"]) == {} and json.loads(h["detail_json"])["overall_only"]
    assert db.judge_human_totals() == [("002_2", 2.0, 5.0)]
    assert db.judge_human_pairs() == []          # 축별 보정에는 안 쓰인다
    assert import_scores(db, out) == 0           # 다시 반입해도 변화 없음
